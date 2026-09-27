from flask import Flask, request, redirect, url_for, render_template_string, send_from_directory
import sqlite3
import os
import html
import webbrowser

app = Flask(__name__)

DB = "vlsi_tracker.db"
PDF_FOLDER = "pdfs"

os.makedirs(PDF_FOLDER, exist_ok=True)


# ---------- DATABASE ----------

def get_db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con


def setup_db():
    con = get_db()

    con.execute("""
        CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            topic TEXT,
            question TEXT,
            answer TEXT,
            status TEXT DEFAULT 'Remaining'
        )
    """)

    # Add status column if old database doesn't have it
    cols = [r["name"] for r in con.execute("PRAGMA table_info(questions)").fetchall()]

    if "status" not in cols:
        try:
            con.execute(
                "ALTER TABLE questions ADD COLUMN status TEXT DEFAULT 'Remaining'"
            )
        except sqlite3.OperationalError:
            pass

    con.commit()
    con.close()


setup_db()


# ---------- HELPERS ----------

def esc(value):
    return html.escape(str(value or ""))


def stats():
    con = get_db()

    total = con.execute(
        "SELECT COUNT(*) FROM questions"
    ).fetchone()[0]

    completed = con.execute(
        "SELECT COUNT(*) FROM questions WHERE status='Completed'"
    ).fetchone()[0]

    remaining = total - completed
    percent = round((completed / total) * 100) if total else 0

    con.close()

    return total, completed, remaining, percent


def layout(content, active="Home"):
    menu = [
        ("🏠", "Home", "/"),
        ("➕", "Add Question", "/add"),
        ("🔍", "Search", "/search"),
        ("📚", "View Questions", "/questions"),
        ("📊", "Import Excel", "/import-excel"),
        ("📄", "Import PDF", "/import-pdf"),
        ("📁", "PDF Library", "/pdf"),
        ("▶️", "YouTube Learning", "/youtube"),
        ("✏️", "Update", "/update"),
        ("🗑️", "Delete", "/delete-page"),
        ("📈", "Your Progress", "/progress"),
        ("ℹ️", "About", "/about")
    ]

    sidebar = ""

    for icon, name, link in menu:
        active_class = "active" if name == active else ""
        sidebar += f"""
        <a class="menu {active_class}"
           href="{link}"
           onclick="closeMenu()">
           <span>{icon}</span>
           <span>{name}</span>
        </a>
        """

    total, completed, remaining, percent = stats()

    return f"""
<!DOCTYPE html>
<html>
<head>

<meta name="viewport"
      content="width=device-width, initial-scale=1, maximum-scale=1">

<title>VLSI Interview Tracker</title>

<style>

* {{
    box-sizing:border-box;
}}

body {{
    margin:0;
    font-family:Arial,Helvetica,sans-serif;
    background:#f5f3fb;
    color:#191536;
}}

a {{
    text-decoration:none;
}}

.sidebar {{
    position:fixed;
    left:0;
    top:0;
    bottom:0;
    width:214px;
    background:#24104f;
    color:white;
    padding:18px 13px;
    z-index:20;
    overflow-y:auto;
}}

.brand {{
    font-size:20px;
    font-weight:bold;
    padding:0 10px 20px;
}}

.brand-small {{
    font-size:11px;
    opacity:.75;
    letter-spacing:1px;
}}

.menu {{
    display:flex;
    gap:10px;
    align-items:center;
    color:white;
    padding:12px 10px;
    margin:4px 0;
    border-radius:9px;
    font-size:14px;
}}

.menu:hover,
.menu.active {{
    background:#43237d;
}}

.main {{
    margin-left:214px;
    min-height:100vh;
}}

.topbar {{
    height:60px;
    background:#24104f;
    color:white;
    display:flex;
    align-items:center;
    padding:0 24px;
}}

.top-title {{
    font-size:24px;
    font-weight:bold;
}}

.hamburger {{
    display:none;
    border:0;
    background:transparent;
    color:white;
    font-size:27px;
    margin-right:12px;
}}

.content {{
    max-width:1200px;
    margin:auto;
    padding:28px;
}}

.hero {{
    background:white;
    border-radius:18px;
    padding:28px 30px;
    margin-bottom:18px;
    box-shadow:0 4px 18px #0000000d;
}}

.hero h1 {{
    margin:0 0 7px;
    font-size:30px;
}}

.hero p {{
    margin:0;
    color:#766f8f;
}}

.cards {{
    display:grid;
    grid-template-columns:repeat(3,1fr);
    gap:14px;
}}

.card {{
    background:white;
    border-radius:16px;
    padding:20px;
    box-shadow:0 4px 16px #0000000c;
}}

.action {{
    min-height:155px;
}}

.icon {{
    width:46px;
    height:46px;
    border-radius:11px;
    display:flex;
    align-items:center;
    justify-content:center;
    font-size:24px;
    margin-bottom:13px;
}}

.green {{
    background:#dff8ed;
}}

.blue {{
    background:#dcecff;
}}

.purple {{
    background:#eadffd;
}}

.red {{
    background:#ffe0e7;
}}

.orange {{
    background:#ffe9ca;
}}

.teal {{
    background:#d9f5ef;
}}

.action h3 {{
    margin:0 0 12px;
    font-size:18px;
}}

.btn {{
    display:inline-block;
    border:0;
    border-radius:8px;
    padding:9px 16px;
    color:white;
    background:#5b35a6;
    cursor:pointer;
}}

.btn.green-btn {{
    background:#19a979;
}}

.btn.blue-btn {{
    background:#287be0;
}}

.btn.red-btn {{
    background:#e54863;
}}

.btn.orange-btn {{
    background:#f39a0b;
}}

.btn.teal-btn {{
    background:#18a67e;
}}

.stats {{
    display:grid;
    grid-template-columns:repeat(4,1fr);
    gap:14px;
    margin-top:18px;
}}

.stat {{
    background:white;
    padding:20px;
    border-radius:16px;
    text-align:center;
    box-shadow:0 4px 16px #0000000c;
}}

.stat-number {{
    font-size:24px;
    font-weight:bold;
    color:#693bb3;
}}

.stat-label {{
    margin-top:5px;
    font-size:14px;
}}

.form-card {{
    background:white;
    padding:25px;
    border-radius:18px;
    box-shadow:0 4px 18px #0000000d;
}}

input, textarea, select {{
    width:100%;
    padding:13px;
    border:1px solid #ddd8e8;
    border-radius:10px;
    margin:7px 0 15px;
    font-size:15px;
}}

textarea {{
    min-height:130px;
    resize:vertical;
}}

label {{
    font-weight:bold;
    font-size:14px;
}}

.question {{
    background:white;
    border-radius:16px;
    padding:18px;
    margin-bottom:13px;
    box-shadow:0 4px 15px #0000000c;
}}

.topic {{
    display:inline-block;
    padding:5px 9px;
    border-radius:20px;
    background:#eee5ff;
    color:#5b35a6;
    font-size:12px;
}}

.answer {{
    background:#f7f6fb;
    padding:12px;
    border-radius:10px;
    margin-top:10px;
    white-space:pre-wrap;
}}

.actions {{
    display:flex;
    gap:8px;
    flex-wrap:wrap;
    margin-top:12px;
}}

.progress-box {{
    background:white;
    padding:25px;
    border-radius:18px;
    box-shadow:0 4px 18px #0000000d;
}}

.progress-track {{
    height:22px;
    background:#e8e5f0;
    border-radius:20px;
    overflow:hidden;
}}

.progress-bar {{
    height:100%;
    background:#693bb3;
    width:{percent}%;
}}

.video {{
    display:flex;
    align-items:center;
    justify-content:space-between;
    gap:15px;
    background:white;
    padding:18px;
    border-radius:15px;
    margin-bottom:12px;
    box-shadow:0 4px 15px #0000000c;
}}

.pdf {{
    display:flex;
    align-items:center;
    justify-content:space-between;
    gap:10px;
    background:white;
    padding:16px;
    border-radius:14px;
    margin-bottom:10px;
}}

.overlay {{
    display:none;
}}

.splash {{
    position:fixed;
    inset:0;
    background:#24104f;
    z-index:9999;
    display:flex;
    justify-content:center;
    align-items:center;
    color:white;
    animation:splashHide 5s forwards;
}}

.splash-inner {{
    text-align:center;
}}

.splash-vlsi {{
    font-size:75px;
    font-weight:bold;
    animation:pulse 1s infinite alternate;
}}

.splash-sub {{
    font-size:16px;
    letter-spacing:3px;
}}

@keyframes pulse {{
    from {{transform:scale(1)}}
    to {{transform:scale(1.08)}}
}}

@keyframes splashHide {{
    0%,88% {{opacity:1; visibility:visible}}
    100% {{opacity:0; visibility:hidden}}
}}

@media(max-width:800px) {{

    .sidebar {{
        transform:translateX(-100%);
        transition:.25s;
        width:250px;
    }}

    .sidebar.open {{
        transform:translateX(0);
    }}

    .main {{
        margin-left:0;
    }}

    .hamburger {{
        display:block;
    }}

    .topbar {{
        padding:0 15px;
        height:58px;
    }}

    .top-title {{
        font-size:20px;
    }}

    .content {{
        padding:14px;
    }}

    .cards {{
    grid-template-columns:repeat(2,1fr);
}}

    }}

    .stats {{
        grid-template-columns:repeat(2,1fr);
    }}

    .hero {{
        padding:22px;
    }}

    .hero h1 {{
        font-size:23px;
    }}

    .overlay.show {{
        display:block;
        position:fixed;
        inset:0;
        background:#0008;
        z-index:15;
    }}
}}

@media(max-width:430px) {{

    .stats {{
        grid-template-columns:1fr 1fr;
        gap:9px;
    }}

    .stat {{
        padding:14px 8px;
    }}

    .stat-number {{
        font-size:21px;
    }}

    .action {{
        min-height:145px;
    }}

}}

</style>

<script>
function openMenu(){{
    document.getElementById("sidebar").classList.add("open");
    document.getElementById("overlay").classList.add("show");
}}

function closeMenu(){{
    document.getElementById("sidebar").classList.remove("open");
    document.getElementById("overlay").classList.remove("show");
}}
</script>

</head>

<body>

<div class="splash">
    <div class="splash-inner">
        <div class="splash-vlsi">VLSI</div>
        <div class="splash-sub">INTERVIEW TRACKER</div>
    </div>
</div>

<div id="overlay" class="overlay" onclick="closeMenu()"></div>

<aside id="sidebar" class="sidebar">

    <div class="brand">
        ⚡ VLSI
        <div class="brand-small">INTERVIEW TRACKER</div>
    </div>

    {sidebar}

</aside>

<div class="main">

    <header class="topbar">
        <button class="hamburger" onclick="openMenu()">☰</button>
        <div class="top-title">VLSI Interview Tracker</div>
    </header>

    <main class="content">
        {content}
    </main>

</div>

</body>
</html>
"""


# ---------- HOME ----------

@app.route("/")
def home():

    total, completed, remaining, percent = stats()

    recent_con = get_db()
    recent = recent_con.execute(
        "SELECT * FROM questions ORDER BY id DESC LIMIT 5"
    ).fetchall()
    recent_con.close()

    recent_html = ""

    for r in recent:
        recent_html += f"""
        <div class="question">
            <span class="topic">{esc(r['topic'])}</span>
            <h3>{esc(r['question'])}</h3>
            <div class="answer">
                {esc(r['answer']) or 'No answer'}
            </div>
        </div>
        """

    body = f"""

    <div class="hero">
        <h1>Welcome to VLSI Interview Tracker 👋</h1>
        <p>Learn • Practice • Prepare • Get Hired</p>
    </div>

    <div class="cards">

        <div class="card action">
            <div class="icon green">➕</div>
            <h3>Add Question</h3>
            <a class="btn green-btn" href="/add">Open</a>
        </div>

        <div class="card action">
            <div class="icon blue">🔍</div>
            <h3>Search</h3>
            <a class="btn blue-btn" href="/search">Open</a>
        </div>

        <div class="card action">
            <div class="icon purple">📚</div>
            <h3>View Questions</h3>
            <a class="btn" href="/questions">Open</a>
        </div>

        <div class="card action">
            <div class="icon red">📁</div>
            <h3>PDF Library</h3>
            <a class="btn red-btn" href="/pdf">Open</a>
        </div>

        <div class="card action">
            <div class="icon orange">▶️</div>
            <h3>YouTube Learning</h3>
            <a class="btn orange-btn" href="/youtube">Open</a>
        </div>

        <div class="card action">
            <div class="icon teal">📈</div>
            <h3>Your Progress</h3>
            <a class="btn teal-btn" href="/progress">Open</a>
        </div>

    </div>

    <div class="stats">

        <div class="stat">
            <div class="stat-number">{total}</div>
            <div class="stat-label">Total Questions</div>
        </div>

        <div class="stat">
            <div class="stat-number">{completed}</div>
            <div class="stat-label">Completed</div>
        </div>

        <div class="stat">
            <div class="stat-number">{remaining}</div>
            <div class="stat-label">Remaining</div>
        </div>

        <div class="stat">
            <div class="stat-number">{percent}%</div>
            <div class="stat-label">Progress</div>
        </div>

    </div>

    <br>

    <div class="hero">
        <h2>Recent Questions</h2>
        {recent_html if recent_html else '<p>No questions yet.</p>'}
    </div>
    """

    return render_template_string(
        layout(body, "Home")
    )


# ---------- ADD ----------

@app.route("/add", methods=["GET", "POST"])
def add_question():

    if request.method == "POST":

        topic = request.form.get("topic", "").strip()
        question = request.form.get("question", "").strip()
        answer = request.form.get("answer", "").strip()

        if topic and question:

            con = get_db()

            con.execute("""
                INSERT INTO questions(topic, question, answer, status)
                VALUES(?,?,?,?)
            """, (topic, question, answer, "Remaining"))

            con.commit()
            con.close()

        return redirect("/questions")

    body = """
    <h2>➕ Add Question</h2>

    <div class="form-card">

    <form method="post">

        <label>Topic</label>
        <input name="topic"
               placeholder="Example: Verilog"
               required>

        <label>Question</label>
        <textarea name="question"
                  placeholder="Enter interview question"
                  required></textarea>

        <label>Answer</label>
        <textarea name="answer"
                  placeholder="Enter answer"></textarea>

        <button class="btn green-btn">
            Save Question
        </button>

    </form>

    </div>
    """

    return render_template_string(layout(body, "Add Question"))


# ---------- VIEW ----------

@app.route("/questions")
def questions():

    con = get_db()

    rows = con.execute("""
        SELECT * FROM questions
        ORDER BY id DESC
    """).fetchall()

    con.close()

    body = "<h2>📚 View Questions</h2>"

    if not rows:
        body += """
        <div class="card">
            <p>No questions found.</p>
        </div>
        """

    for r in rows:

        status = r["status"] or "Remaining"

        body += f"""
        <div class="question">

            <span class="topic">
                {esc(r['topic'])}
            </span>

            <h3>
                {esc(r['question'])}
            </h3>

            <div class="answer">
                <b>Answer:</b><br>
                {esc(r['answer']) or 'No answer added'}
            </div>

            <p>
                <b>Status:</b> {esc(status)}
            </p>

            <div class="actions">

                <a class="btn"
                   href="/edit/{r['id']}">
                   ✏️ Update
                </a>

                <a class="btn"
                   style="background:#e54863"
                   href="/delete/{r['id']}"
                   onclick="return confirm('Delete this question?')">
                   🗑️ Delete
                </a>

            </div>

        </div>
        """

    return render_template_string(layout(body, "View Questions"))


# ---------- EDIT ----------

@app.route("/edit/<int:qid>", methods=["GET", "POST"])
def edit(qid):

    con = get_db()

    if request.method == "POST":

        topic = request.form.get("topic", "").strip()
        question = request.form.get("question", "").strip()
        answer = request.form.get("answer", "").strip()
        status = request.form.get("status", "Remaining")

        con.execute("""
            UPDATE questions
            SET topic=?, question=?, answer=?, status=?
            WHERE id=?
        """, (topic, question, answer, status, qid))

        con.commit()
        con.close()

        return redirect("/questions")

    row = con.execute(
        "SELECT * FROM questions WHERE id=?",
        (qid,)
    ).fetchone()

    con.close()

    if not row:
        return redirect("/questions")

    body = f"""

    <h2>✏️ Update Question</h2>

    <div class="form-card">

    <form method="post">

        <label>Topic</label>
        <input name="topic"
               value="{esc(row['topic'])}"
               required>

        <label>Question</label>
        <textarea name="question"
                  required>{esc(row['question'])}</textarea>

        <label>Answer</label>
        <textarea name="answer">{esc(row['answer'])}</textarea>

        <label>Status</label>

        <select name="status">

            <option value="Remaining"
                {"selected" if row["status"] != "Completed" else ""}>
                Remaining
            </option>

            <option value="Completed"
                {"selected" if row["status"] == "Completed" else ""}>
                Completed
            </option>

        </select>

        <button class="btn">
            Update Question
        </button>

    </form>

    </div>
    """

    return render_template_string(layout(body, "Update"))


@app.route("/update")
def update_page():

    body = """

    <h2>✏️ Update Question</h2>

    <div class="card">
        <p>Select a question below to update it.</p>
    </div>

    """

    con = get_db()

    rows = con.execute(
        "SELECT * FROM questions ORDER BY id DESC"
    ).fetchall()

    con.close()

    for r in rows[:50]:

        body += f"""
        <div class="question">

            <span class="topic">
                ID {r['id']} • {esc(r['topic'])}
            </span>

            <h3>{esc(r['question'])}</h3>

            <a class="btn"
               href="/edit/{r['id']}">
               ✏️ Update
            </a>

        </div>
        """

    return render_template_string(layout(body, "Update"))


# ---------- DELETE ----------

@app.route("/delete/<int:qid>")
def delete(qid):

    con = get_db()

    con.execute(
        "DELETE FROM questions WHERE id=?",
        (qid,)
    )

    con.commit()
    con.close()

    return redirect("/questions")


@app.route("/delete-page")
def delete_page():

    body = """
    <h2>🗑️ Delete Question</h2>
    """

    con = get_db()

    rows = con.execute(
        "SELECT * FROM questions ORDER BY id DESC"
    ).fetchall()

    con.close()

    for r in rows[:50]:

        body += f"""
        <div class="question">

            <span class="topic">
                ID {r['id']} • {esc(r['topic'])}
            </span>

            <h3>{esc(r['question'])}</h3>

            <a class="btn"
               style="background:#e54863"
               href="/delete/{r['id']}"
               onclick="return confirm('Delete this question?')">
               🗑️ Delete
            </a>

        </div>
        """

    return render_template_string(layout(body, "Delete"))


# ---------- SEARCH ----------

@app.route("/search", methods=["GET", "POST"])
def search():

    results = ""

    if request.method == "POST":

        text = request.form.get("text", "").strip()

        con = get_db()

        rows = con.execute("""
            SELECT * FROM questions
            WHERE topic LIKE ?
               OR question LIKE ?
               OR answer LIKE ?
            ORDER BY id DESC
        """, (
            "%" + text + "%",
            "%" + text + "%",
            "%" + text + "%"
        )).fetchall()

        con.close()

        for r in rows:

            results += f"""
            <div class="question">

                <span class="topic">
                    {esc(r['topic'])}
                </span>

                <h3>{esc(r['question'])}</h3>

                <div class="answer">
                    {esc(r['answer']) or 'No answer'}
                </div>

            </div>
            """

        if not rows:
            results = """
            <div class="card">
                <p>No results found.</p>
            </div>
            """

    body = f"""

    <h2>🔍 Search</h2>

    <div class="form-card">

        <form method="post">

            <input name="text"
                   placeholder="Search topic, question or answer"
                   required>

            <button class="btn blue-btn">
                🔍 Search
            </button>

        </form>

    </div>

    <br>

    {results}

    """

    return render_template_string(layout(body, "Search"))


# ---------- PROGRESS ----------

@app.route("/progress")
def progress():

    total, completed, remaining, percent = stats()

    body = f"""

    <h2>📈 Your Progress</h2>

    <div class="stats">

        <div class="stat">
            <div class="stat-number">{total}</div>
            <div class="stat-label">Total Questions</div>
        </div>

        <div class="stat">
            <div class="stat-number">{completed}</div>
            <div class="stat-label">Completed</div>
        </div>

        <div class="stat">
            <div class="stat-number">{remaining}</div>
            <div class="stat-label">Remaining</div>
        </div>

        <div class="stat">
            <div class="stat-number">{percent}%</div>
            <div class="stat-label">Progress</div>
        </div>

    </div>

    <br>

    <div class="progress-box">

        <h3>Preparation Progress</h3>

        <div class="progress-track">
            <div class="progress-bar"></div>
        </div>

        <p>
            {completed} of {total} questions completed.
        </p>

    </div>

    <br>

    <div class="hero">
        <h3>💡 Keep going!</h3>
        <p>
            Every question you practice brings you
            closer to your goal.
        </p>
    </div>
    """

    return render_template_string(layout(body, "Your Progress"))


# ---------- YOUTUBE ----------

@app.route("/youtube")
def youtube():

    body = """

    <h2>▶️ YouTube Learning</h2>

    <div class="video">
        <div>
            <h3>📺 Neso Academy</h3>
            <p>Digital Electronics and VLSI basics</p>
        </div>

        <a class="btn orange-btn"
           target="_blank"
           href="https://www.youtube.com/@nesoacademy">
           Watch
        </a>
    </div>

    <div class="video">
        <div>
            <h3>💻 SystemVerilog</h3>
            <p>SystemVerilog tutorials</p>
        </div>

        <a class="btn blue-btn"
           target="_blank"
           href="https://www.youtube.com/results?search_query=SystemVerilog+tutorial">
           Watch
        </a>
    </div>

    <div class="video">
        <div>
            <h3>🔧 UVM</h3>
            <p>UVM verification tutorials</p>
        </div>

        <a class="btn green-btn"
           target="_blank"
           href="https://www.youtube.com/results?search_query=UVM+tutorial">
           Watch
        </a>
    </div>

    <div class="video">
        <div>
            <h3>🎯 VLSI Interview</h3>
            <p>Interview preparation questions</p>
        </div>

        <a class="btn red-btn"
           target="_blank"
           href="https://www.youtube.com/results?search_query=VLSI+interview+questions">
           Watch
        </a>
    </div>

    """

    return render_template_string(layout(body, "YouTube Learning"))


# ---------- PDF LIBRARY ----------

@app.route("/pdf")
def pdf_library():

    files = []

    if os.path.exists(PDF_FOLDER):
        files = [
            f for f in os.listdir(PDF_FOLDER)
            if f.lower().endswith(".pdf")
        ]

    body = """
    <h2>📁 PDF Library</h2>

    <div class="card">
        <a class="btn red-btn" href="/import-pdf">
            ➕ Add PDF
        </a>
    </div>

    <br>
    """

    if not files:

        body += """
        <div class="card">
            <p>No PDF files added yet.</p>
        </div>
        """

    for filename in files:

        safe_name = esc(filename)

        body += f"""
        <div class="pdf">

            <div>
                <b>📄 {safe_name}</b>
            </div>

            <div>
                <a class="btn blue-btn"
                   href="/pdf-file/{safe_name}"
                   target="_blank">
                   Open
                </a>
            </div>

        </div>
        """

    return render_template_string(layout(body, "PDF Library"))


@app.route("/pdf-file/<path:filename>")
def pdf_file(filename):
    return send_from_directory(PDF_FOLDER, filename)


# ---------- IMPORT EXCEL ----------

@app.route("/import-excel", methods=["GET", "POST"])
def import_excel():

    message = ""

    if request.method == "POST":

        file = request.files.get("file")

        if file and file.filename:

            try:
                from openpyxl import load_workbook

                wb = load_workbook(
                    file,
                    data_only=True
                )

                ws = wb.active

                con = get_db()
                count = 0

                for row in ws.iter_rows(values_only=True):

                    values = list(row)

                    if not values:
                        continue

                    topic = str(values[0] or "").strip()
                    question = str(
                        values[1] if len(values) > 1 else ""
                    ).strip()

                    answer = str(
                        values[2] if len(values) > 2 else ""
                    ).strip()

                    status = str(
                        values[3] if len(values) > 3 else "Remaining"
                    ).strip()

                    if topic.lower() == "topic" and question.lower() == "question":
                        continue

                    if topic and question:

                        if status not in ("Completed", "Remaining"):
                            status = "Remaining"

                        con.execute("""
                            INSERT INTO questions
                            (topic,question,answer,status)
                            VALUES(?,?,?,?)
                        """, (
                            topic,
                            question,
                            answer,
                            status
                        ))

                        count += 1

                con.commit()
                con.close()

                message = f"{count} questions imported successfully."

            except Exception as e:
                message = "Excel import error: " + str(e)

    body = f"""

    <h2>📊 Import Excel</h2>

    <div class="form-card">

        <form method="post"
              enctype="multipart/form-data">

            <label>Select Excel File</label>

            <input type="file"
                   name="file"
                   accept=".xlsx,.xls"
                   required>

            <button class="btn green-btn">
                Import Questions
            </button>

        </form>

        <br>

        <p>
        Excel columns:
        <b>Topic | Question | Answer | Status</b>
        </p>

        <p>{esc(message)}</p>

    </div>

    """

    return render_template_string(layout(body, "Import Excel"))


# ---------- IMPORT PDF ----------

@app.route("/import-pdf", methods=["GET", "POST"])
def import_pdf():

    message = ""

    if request.method == "POST":

        file = request.files.get("file")

        if file and file.filename:

            filename = os.path.basename(file.filename)

            if not filename.lower().endswith(".pdf"):
                message = "Please select a PDF file."

            else:

                save_path = os.path.join(
                    PDF_FOLDER,
                    filename
                )

                file.save(save_path)

                message = "PDF added to PDF Library."

    body = f"""

    <h2>📄 Import PDF</h2>

    <div class="form-card">

        <form method="post"
              enctype="multipart/form-data">

            <label>Select PDF</label>

            <input type="file"
                   name="file"
                   accept=".pdf"
                   required>

            <button class="btn red-btn">
                Add PDF
            </button>

        </form>

        <br>

        <p>{esc(message)}</p>

    </div>

    """

    return render_template_string(layout(body, "Import PDF"))


# ---------- ABOUT ----------

@app.route("/about")
def about():

    body = """

    <div class="hero">
        <h1>⚡ VLSI Interview Tracker</h1>
        <p>
            Learn • Practice • Prepare • Get Hired
        </p>
    </div>

    <div class="card">

        <h2>About</h2>

        <p>
        VLSI Interview Tracker helps students
        organize interview questions and track
        their preparation.
        </p>

        <p>➕ Add Questions</p>
        <p>🔍 Search Questions</p>
        <p>📚 View Questions</p>
        <p>✏️ Update Questions</p>
        <p>🗑️ Delete Questions</p>
        <p>📊 Import Excel</p>
        <p>📄 Import PDF</p>
        <p>📁 PDF Library</p>
        <p>▶️ YouTube Learning</p>
        <p>📈 Progress Tracking</p>

    </div>

    """

    return render_template_string(layout(body, "About"))


# ---------- START ----------

if __name__ == "__main__":

    print()
    print("===================================")
    print(" VLSI INTERVIEW TRACKER MOBILE")
    print("===================================")
    print("Open on laptop:")
    print("http://127.0.0.1:5000")
    print()
    print("Open on phone:")
    print("http://192.168.55.103:5000")
    print()

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )